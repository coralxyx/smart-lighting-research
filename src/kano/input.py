from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


REQUIRED_COLUMNS = (
    "原句 (未分句)",
    "分句 (模型输入)",
    "使用场景分类",
    "标准美学指标分类",
    "用户感受分类",
    "情感强度",
)


@dataclass(frozen=True)
class TaggedRecord:
    row_number: int
    full_text: str
    sentence: str
    scene: str
    aesthetic: str
    perception: str
    sentiment: str


def _clean(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("\r", "\n").split())


def load_tagged_records(path: str | Path) -> list[TaggedRecord]:
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    worksheet = workbook.active
    headers = [_clean(cell.value) for cell in worksheet[1]]
    column_index = {name: index for index, name in enumerate(headers)}

    missing = [name for name in REQUIRED_COLUMNS if name not in column_index]
    if missing:
        raise ValueError(f"Missing required columns: {missing}")

    records: list[TaggedRecord] = []
    for row_number, values in enumerate(
        worksheet.iter_rows(min_row=2, values_only=True), start=2
    ):
        records.append(
            TaggedRecord(
                row_number=row_number,
                full_text=_clean(values[column_index["原句 (未分句)"]]),
                sentence=_clean(values[column_index["分句 (模型输入)"]]),
                scene=_clean(values[column_index["使用场景分类"]]),
                aesthetic=_clean(values[column_index["标准美学指标分类"]]),
                perception=_clean(values[column_index["用户感受分类"]]),
                sentiment=_clean(values[column_index["情感强度"]]),
            )
        )

    workbook.close()
    return records


def validate_records(records: Iterable[TaggedRecord]) -> dict[str, object]:
    records = list(records)
    required_fields = ("full_text", "scene", "aesthetic", "perception", "sentiment")
    missing_by_field = {
        field: sum(not getattr(record, field) for record in records)
        for field in required_fields
    }
    invalid_sentiment_rows = [
        record.row_number
        for record in records
        if record.sentiment not in {"正向", "中性", "负向"}
    ]
    return {
        "missing_by_field": missing_by_field,
        "invalid_sentiment_rows": invalid_sentiment_rows,
    }
