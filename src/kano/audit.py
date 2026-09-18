from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


def _reference_global_counts(reference_path: Path) -> dict[str, object]:
    workbook = load_workbook(reference_path, read_only=False, data_only=True)
    worksheet = workbook["全局统计"]
    result = {
        "record_count": worksheet["B4"].value,
        "scene_counts": {
            worksheet.cell(row, 1).value: worksheet.cell(row, 2).value
            for row in range(10, 17)
            if worksheet.cell(row, 1).value
        },
        "aesthetic_counts": {
            worksheet.cell(row, 5).value: worksheet.cell(row, 7).value
            for row in range(10, 19)
            if worksheet.cell(row, 5).value
        },
        "perception_counts": {
            worksheet.cell(row, 1).value: worksheet.cell(row, 2).value
            for row in range(23, 30)
            if worksheet.cell(row, 1).value
        },
        "sentiment_counts": {
            worksheet.cell(row, 5).value: worksheet.cell(row, 6).value
            for row in range(23, 26)
            if worksheet.cell(row, 5).value
        },
    }
    workbook.close()
    return result


def _reference_scene_aesthetic(reference_path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(reference_path, read_only=False, data_only=True)
    worksheet = workbook["Kano汇总"]
    rows: list[dict[str, object]] = []
    for row in range(4, worksheet.max_row + 1):
        count = worksheet.cell(row, 4).value
        if not isinstance(count, (int, float)):
            continue
        rows.append(
            {
                "scene": worksheet.cell(row, 1).value,
                "aesthetic": worksheet.cell(row, 3).value,
                "count": int(count),
                "positive_count": int(worksheet.cell(row, 6).value or 0),
                "neutral_count": int(worksheet.cell(row, 7).value or 0),
                "negative_count": int(worksheet.cell(row, 8).value or 0),
                "legacy_classification": worksheet.cell(row, 14).value,
            }
        )
    workbook.close()
    return rows


def _as_count_dict(rows: list[dict[str, object]], key: str) -> dict[str, int]:
    return {str(row[key]): int(row["count"]) for row in rows}


def _compare_count_dicts(
    computed: dict[str, int], reference: dict[str, int]
) -> list[dict[str, object]]:
    differences = []
    for key in sorted(set(computed) | set(reference)):
        computed_value = computed.get(key, 0)
        reference_value = reference.get(key, 0)
        if computed_value != reference_value:
            differences.append(
                {
                    "key": key,
                    "computed": computed_value,
                    "reference": reference_value,
                    "difference": computed_value - reference_value,
                }
            )
    return differences


def audit_against_reference(
    aggregates: dict[str, object], reference_path: str | Path
) -> dict[str, object]:
    reference_path = Path(reference_path)
    reference_global = _reference_global_counts(reference_path)
    computed_global = {
        "scene_counts": _as_count_dict(aggregates["scene_counts"], "scene"),
        "aesthetic_counts": _as_count_dict(
            aggregates["aesthetic_counts"], "aesthetic"
        ),
        "perception_counts": _as_count_dict(
            aggregates["perception_counts"], "perception"
        ),
        "sentiment_counts": _as_count_dict(
            aggregates["sentiment_counts"], "sentiment"
        ),
    }
    global_differences = {
        name: _compare_count_dicts(computed_global[name], reference_global[name])
        for name in computed_global
    }

    computed_scene_rows = {
        (str(row["scene"]), str(row["aesthetic"])): row
        for row in aggregates["scene_aesthetic"]
        if row["scene"] != "无"
    }
    reference_scene_rows = {
        (str(row["scene"]), str(row["aesthetic"])): row
        for row in _reference_scene_aesthetic(reference_path)
    }
    count_fields = ("count", "positive_count", "neutral_count", "negative_count")
    scene_differences: list[dict[str, object]] = []
    for key in sorted(set(computed_scene_rows) | set(reference_scene_rows)):
        computed_row = computed_scene_rows.get(key)
        reference_row = reference_scene_rows.get(key)
        if computed_row is None or reference_row is None:
            scene_differences.append(
                {
                    "scene": key[0],
                    "aesthetic": key[1],
                    "reason": "missing group",
                }
            )
            continue
        differences = {
            field: int(computed_row[field]) - int(reference_row[field])
            for field in count_fields
            if int(computed_row[field]) != int(reference_row[field])
        }
        if differences:
            scene_differences.append(
                {"scene": key[0], "aesthetic": key[1], "differences": differences}
            )

    return {
        "reference_path": str(reference_path),
        "record_count": {
            "computed": aggregates["record_count"],
            "reference": reference_global["record_count"],
            "match": aggregates["record_count"] == reference_global["record_count"],
        },
        "global_differences": global_differences,
        "scene_aesthetic_differences": scene_differences,
        "checks": {
            "global_counts_match": not any(global_differences.values()),
            "scene_aesthetic_counts_match": not scene_differences,
        },
        "legacy_classification_counts": dict(
            Counter(
                row["legacy_classification"] for row in reference_scene_rows.values()
            )
        ),
        "classification_rule_status": (
            "Legacy classifications are static cell values; executable thresholds "
            "were not found in the workbook."
        ),
    }
