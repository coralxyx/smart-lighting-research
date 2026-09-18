from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


SCENE_SHEET_PREFIX = "Kano_"


def load_legacy_metadata(reference_path: str | Path) -> dict[str, object]:
    """Read the non-recoverable labels and display metadata from the old workbook."""
    workbook = load_workbook(Path(reference_path), read_only=False, data_only=True)

    method_sheet = workbook["方法说明"]
    touchpoint_by_aesthetic = {
        str(method_sheet.cell(row, 4).value): str(method_sheet.cell(row, 5).value)
        for row in range(4, method_sheet.max_row + 1)
        if method_sheet.cell(row, 4).value and method_sheet.cell(row, 5).value
    }

    scene_aesthetic: dict[tuple[str, str], dict[str, str]] = {}
    for worksheet in workbook.worksheets:
        if not worksheet.title.startswith(SCENE_SHEET_PREFIX):
            continue
        scene = worksheet.title[len(SCENE_SHEET_PREFIX) :]
        for row in range(5, worksheet.max_row + 1):
            aesthetic = worksheet.cell(row, 2).value
            if not aesthetic:
                continue
            scene_aesthetic[(scene, str(aesthetic))] = {
                "touchpoint": str(worksheet.cell(row, 1).value or ""),
                "classification": str(worksheet.cell(row, 13).value or ""),
                "priority": str(worksheet.cell(row, 14).value or ""),
                "recommendation": str(worksheet.cell(row, 15).value or ""),
            }

    workbook.close()
    return {
        "touchpoint_by_aesthetic": touchpoint_by_aesthetic,
        "scene_aesthetic": scene_aesthetic,
    }
