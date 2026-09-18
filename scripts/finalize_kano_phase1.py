from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
PAYLOAD_PATH = PROJECT_ROOT / "analysis" / "generated" / "kano-phase1-workbook-data.json"
REFERENCE_PATH = PROJECT_ROOT / "analysis" / "kano-pair-statistics.xlsx"
OUTPUT_PATH = PROJECT_ROOT / "outputs" / "kano_phase1" / "kano-phase1-reproduced.xlsx"


def load_legacy_row_order() -> dict[tuple[str, str], int]:
    workbook = load_workbook(REFERENCE_PATH, read_only=False, data_only=True)
    worksheet = workbook["Kano汇总"]
    order = {
        (str(worksheet.cell(row, 1).value), str(worksheet.cell(row, 3).value)): row
        for row in range(4, worksheet.max_row + 1)
        if worksheet.cell(row, 1).value and worksheet.cell(row, 3).value
    }
    workbook.close()
    return order


def main() -> None:
    subprocess.run(
        [sys.executable, str(PROJECT_ROOT / "scripts" / "run_kano_phase1.py")],
        cwd=PROJECT_ROOT,
        check=True,
    )

    payload = json.loads(PAYLOAD_PATH.read_text(encoding="utf-8"))
    legacy_order = load_legacy_row_order()
    payload["summary_rows"].sort(
        key=lambda row: legacy_order[(str(row["scene"]), str(row["aesthetic"]))]
    )
    PAYLOAD_PATH.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    subprocess.run(
        [
            "node",
            str(PROJECT_ROOT / "scripts" / "build_kano_workbook.mjs"),
            str(PAYLOAD_PATH),
            str(OUTPUT_PATH),
        ],
        cwd=PROJECT_ROOT,
        check=True,
    )
    print(f"Final workbook: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
