from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


def equal(left: object, right: object) -> bool:
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        return abs(float(left) - float(right)) < 1e-10
    return left == right


def compare(reference_path: Path, reproduced_path: Path) -> dict[str, object]:
    reference = load_workbook(reference_path, data_only=True)
    reproduced = load_workbook(reproduced_path, data_only=True)
    specifications = [("Kano汇总", 4, 1, 15), ("词对明细", 4, 1, 13)]
    specifications.extend(
        (name, 5, 1, 15)
        for name in reference.sheetnames
        if name.startswith("Kano_")
    )

    result: dict[str, object] = {}
    for sheet_name, start_row, first_column, last_column in specifications:
        old_sheet = reference[sheet_name]
        new_sheet = reproduced[sheet_name]
        old_rows = [
            [
                old_sheet.cell(row, column).value
                for column in range(first_column, last_column + 1)
            ]
            for row in range(start_row, old_sheet.max_row + 1)
        ]
        new_rows = [
            [
                new_sheet.cell(row, column).value
                for column in range(first_column, last_column + 1)
            ]
            for row in range(start_row, new_sheet.max_row + 1)
        ]
        differences = sum(
            not equal(old_value, new_value)
            for old_row, new_row in zip(old_rows, new_rows)
            for old_value, new_value in zip(old_row, new_row)
        )
        differences += abs(len(old_rows) - len(new_rows))
        result[sheet_name] = {
            "reference_rows": len(old_rows),
            "reproduced_rows": len(new_rows),
            "cell_differences": differences,
        }

    reference.close()
    reproduced.close()
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description="Compare phase-one Kano workbooks.")
    parser.add_argument("reference", type=Path)
    parser.add_argument("reproduced", type=Path)
    args = parser.parse_args()
    result = compare(args.reference, args.reproduced)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if any(item["cell_differences"] for item in result.values()):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
